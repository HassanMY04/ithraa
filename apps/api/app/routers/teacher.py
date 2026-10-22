"""Teacher screens: dashboard, session planner, gradebook — and every action on them."""
import io
import random
from collections import defaultdict
from datetime import datetime, timedelta

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select

from app.core import clock
from app.core.deps import DB, TeacherUser
from app.models import (
    Assignment,
    Attendance,
    Group,
    GroupMember,
    Lesson,
    Level,
    LibraryItem,
    Material,
    Note,
    PortfolioItem,
    Question,
    Recommendation,
    Skill,
    SpecTrack,
    StudentProfile,
    StudentSkill,
    Submission,
    Tip,
    Track,
    User,
)
from app.services import events
from app.services import lessons as lesson_svc
from app.services import teacher_automations as autos
from app.services import teachers as teacher_svc
from app.services.files import kind_for, save_upload, size_label
from app.services.metrics import group_member_ids, load_stats
from app.services.progress import set_skill
from app.services.schedule import LessonTime, Planner

router = APIRouter(prefix="/teacher", tags=["teacher"])

# Level track -> specialization track the teacher can recommend for
REC_TRACK = {"cyber": "cyber", "ai": "ai", "design": "design", "coding": "software"}
# Specialization -> general track used for portfolio items
SPEC_TO_TRACK = {"cyber": "cyber", "ai": "ai", "software": "coding", "design": "design", "robotics": "coding"}
BLOCK_KINDS = {"talk", "teach", "act", "quiz", "wrap"}


# ------------------------------------------------------------------ helpers
def own_group(db, teacher: User, group_id: int) -> Group:
    g = db.get(Group, group_id)
    if not g or g.teacher_id != teacher.id:
        raise HTTPException(404, "Group not found")
    return g


def own_lesson(db, teacher: User, lesson_id: int) -> Lesson:
    lesson = db.get(Lesson, lesson_id)
    if not lesson:
        raise HTTPException(404, "Session not found")
    own_group(db, teacher, lesson.group_id)
    return lesson


def own_student(db, teacher: User, student_id: int) -> StudentProfile:
    ok = db.scalar(
        select(GroupMember.id).join(Group, Group.id == GroupMember.group_id)
        .where(GroupMember.student_id == student_id, Group.teacher_id == teacher.id)
    )
    if not ok:
        raise HTTPException(404, "Not one of your students")
    return db.get(StudentProfile, student_id)


def short(g: Group) -> str:
    return teacher_svc.group_short(g)


def classes_of(db, ids: list[int]) -> list[str]:
    return sorted({p.class_name for p in db.scalars(select(StudentProfile).where(StudentProfile.id.in_(ids or [0])))})


def day_word(d) -> str:
    diff = (d - clock.today()).days
    if diff == 0:
        return "Today"
    if diff == 1:
        return "Tomorrow"
    return clock.fmt_day(d)


def readiness(db, lesson: Lesson, upcoming: bool = True) -> list[dict]:
    items = lesson_svc.readiness(db, lesson, upcoming)
    sm = lesson_svc.slides_material(lesson)
    for it in items:
        if it["key"] == "slides" and sm and not lesson.slides and "slides" in sm.meta:
            it["m"] = f"{sm.meta.split(' · ')[0]} · shared with students"
    return items


def teacher_students(db, teacher: User) -> list[int]:
    return list(db.scalars(
        select(func.distinct(GroupMember.student_id)).join(Group, Group.id == GroupMember.group_id)
        .where(Group.teacher_id == teacher.id)
    ))


def sub_counts(db, assignment_id: int) -> int:
    return db.scalar(select(func.count()).select_from(Submission).where(
        Submission.assignment_id == assignment_id, Submission.status.in_(["submitted", "graded"]))) or 0


def level_of(db, g: Group) -> tuple[Level | None, Track | None]:
    if not g.level_id:
        return None, None
    lv = db.get(Level, g.level_id)
    return lv, db.get(Track, lv.track_id)


# ---------------------------------------------------------------- dashboard
@router.get("/dashboard")
def dashboard(user: TeacherUser, db: DB) -> dict:
    planner = Planner(db)
    groups = teacher_svc.teacher_groups(db, user.id)
    gids = [g.id for g in groups]
    local = {g.id: load_stats(db, group_member_ids(db, g.id), planner, group_ids=[g.id]) for g in groups}
    autos.run(db, user, planner, local)
    on = autos.settings(db, user.id)
    student_ids = teacher_students(db, user)

    # KPIs
    waiting = db.execute(
        select(Assignment.kind, func.count(Submission.id))
        .join(Submission, Submission.assignment_id == Assignment.id)
        .where(Assignment.group_id.in_(gids or [0]), Submission.status == "submitted")
        .group_by(Assignment.kind)
    ).all()
    wmap = dict(waiting)
    k = teacher_svc.kpis(db, user.id, planner)
    hw_all = []
    on_time_by_group = []
    for g in groups:
        rows_g = [h for s in local[g.id].values() for h in s.homework]
        hw_all += rows_g
        if rows_g:
            on_time_by_group.append((short(g), round(sum(1 for h in rows_g if h.on_time) / len(rows_g) * 100)))
    hw_on = round(sum(1 for h in hw_all if h.on_time) / len(hw_all) * 100) if hw_all else None
    lowest = min(on_time_by_group, key=lambda r: r[1]) if len(on_time_by_group) > 1 else None

    stats = load_stats(db, student_ids, planner)
    flags = []
    if on["flag"]:
        for sid in student_ids:
            s = stats[sid]
            if s.flags():
                u, p = db.get(User, sid), db.get(StudentProfile, sid)
                flags.append({"student_id": sid, "name": u.name, "first": u.first_name, "cls": p.class_name,
                              "why": s.short_reason(), "parent": p.parent_name})
        flags.sort(key=lambda f: (-int(f["cls"].split("-")[0]), f["name"]))

    # Next session
    nxt = lesson_svc.next_lesson(db, planner, gids)
    next_card = None
    roster = None
    if nxt:
        lesson, t = nxt
        g = db.get(Group, lesson.group_id)
        members = group_member_ids(db, g.id)
        ready = readiness(db, lesson, upcoming=True)
        n_ok = sum(1 for r in ready if r["ok"])
        is_today = t.day == clock.today()
        next_card = {
            "lesson_id": lesson.id, "group_id": g.id, "group_name": g.name,
            "label": f"Next session · {day_word(t.day)} {clock.fmt_time(t.start)} · {g.room}",
            "title": f"{short(g)} · Week {lesson_svc.module_week_number(db, g, lesson.week)} · {lesson.title}",
            "classes": f"{len(members)} students from {', '.join(classes_of(db, members))}",
            "ready": ready, "ready_count": n_ok, "ready_total": len(ready),
            "has_activity": lesson_svc.has_activity(lesson), "teams": len(lesson.teams or []),
            "is_today": is_today, "activity_open": lesson.activity_open,
            "checkin_code": autos.ensure_checkin_code(lesson) if on["attendance"] else None,
        }
        last = lesson_svc.last_lesson(db, planner, g)
        if last:
            ll, lt = last
            status = dict(db.execute(select(Attendance.student_id, Attendance.status)
                                     .where(Attendance.lesson_id == ll.id)).all())
            people = []
            for sid in members:
                u, p = db.get(User, sid), db.get(StudentProfile, sid)
                people.append({"id": sid, "name": u.name, "ini": u.initials, "cls": p.class_name,
                               "status": status.get(sid)})
            roster = {"lesson_id": ll.id, "group": short(g), "when": clock.fmt_day(lt.day), "students": people}

    # To grade
    to_grade = []
    rows_tg = db.execute(
        select(Assignment, func.count(Submission.id))
        .join(Submission, Submission.assignment_id == Assignment.id)
        .where(Assignment.group_id.in_(gids or [0]), Submission.status == "submitted")
        .group_by(Assignment.id).order_by(Assignment.due_at)
    ).all()
    for a, n in rows_tg:
        g = db.get(Group, a.group_id)
        due = a.due_at.date()
        if (clock.today() - due).days == 1:
            when = "due yesterday"
        elif due == clock.today():
            when = "due today"
        elif due > clock.today():
            when = f"due {clock.fmt_day(due)}"
        else:
            when = f"due {clock.fmt_day(due)}"
        to_grade.append({"assignment_id": a.id, "group_id": g.id, "t": a.title, "m": f"{short(g)} · {when}", "n": n})

    # Tip audiences: the grades this teacher teaches
    grade_opts = []
    seen = set()
    for g in groups:
        if g.grade_key in seen:
            continue
        seen.add(g.grade_key)
        if int(g.grade_key) >= 10:
            n = len(group_member_ids(db, g.id))
            grade_opts.append({"key": "10", "grades": [10, 11, 12], "label": f"{g.spec_track.short} · {n}", "n": n,
                               "group_id": g.id})
        else:
            gr = int(g.grade_key)
            n = db.scalar(select(func.count()).select_from(StudentProfile).where(StudentProfile.grade == gr)) or 0
            grade_opts.append({"key": g.grade_key, "grades": [gr], "label": f"Grade {gr} · {n}", "n": n})
    default_key = db.get(Group, next_card["group_id"]).grade_key if next_card else (grade_opts[0]["key"] if grade_opts else None)
    sched_day = nxt[1].day if nxt else clock.today() + timedelta(days=1)

    questions = [
        {"id": q.id, "student": db.get(User, q.student_id).name, "cls": db.get(StudentProfile, q.student_id).class_name,
         "context": q.context, "text": q.text, "when": clock.relative(q.created_at), "answer": q.answer}
        for q in db.scalars(select(Question).where(Question.teacher_id == user.id)
                            .order_by(Question.answered_at.is_not(None), Question.created_at.desc()).limit(12))
    ]
    db.commit()
    return {
        "date": f"{clock.fmt_long(clock.today())} · Week {clock.current_week()}",
        "greeting": clock.greeting(), "greeting_ar": clock.greeting("ar"), "first_name": user.first_name,
        "kpis": {
            "to_grade": sum(wmap.values()),
            "to_grade_split": " · ".join(x for x in [
                f"{wmap['journal']} journal{'s' if wmap['journal'] != 1 else ''}" if wmap.get("journal") else "",
                f"{wmap['homework']} homework{'s' if wmap['homework'] != 1 else ''}" if wmap.get("homework") else "",
            ] if x) or "All caught up",
            "attendance": k.attendance_pct, "students": len(student_ids),
            "hw_on_time": hw_on,
            "hw_note": f"{lowest[0]} lowest at {lowest[1]}%" if lowest else "This term",
            "attention": len(flags), "flag_on": on["flag"],
        },
        "next": next_card, "flags": flags, "roster": roster, "to_grade": to_grade,
        "tip": {"options": grade_opts, "default": default_key,
                "schedule_label": f"Schedule for {clock.DAYS_SHORT[sched_day.weekday()]} 8:00",
                "schedule_at": datetime.combine(sched_day, datetime.min.time()).replace(hour=8).isoformat()},
        "questions": questions,
    }


@router.post("/lessons/{lesson_id}/teams")
def auto_teams(lesson_id: int, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    members = group_member_ids(db, lesson.group_id)
    by_class: dict[str, list[int]] = defaultdict(list)
    for sid in members:
        by_class[db.get(StudentProfile, sid).class_name].append(sid)
    for v in by_class.values():
        random.shuffle(v)
    # Deal students round-robin across classes so every team mixes classes.
    order: list[int] = []
    pools = sorted(by_class.values(), key=len, reverse=True)
    while any(pools):
        for p in pools:
            if p:
                order.append(p.pop())
    n_teams = max(1, len(order) // 3)
    teams: list[list[int]] = [[] for _ in range(n_teams)]
    for i, sid in enumerate(order):
        teams[i % n_teams].append(sid)
    lesson.teams = teams
    db.commit()
    return {"teams": len(teams), "size": 3}


class ActivityIn(BaseModel):
    open: bool


@router.post("/lessons/{lesson_id}/activity")
def toggle_activity(lesson_id: int, body: ActivityIn, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    lesson.activity_open = body.open
    if body.open:
        for sid in group_member_ids(db, lesson.group_id):
            events.notify(db, sid, f"{user.formal_short} started the live activity — open your session page.",
                          link=f"/student/session/{lesson.id}", ref=f"activity:{lesson.id}")
    db.commit()
    return {"open": lesson.activity_open}


class AttendanceIn(BaseModel):
    student_id: int
    status: str = Field(pattern="^[PLA]$")


@router.put("/lessons/{lesson_id}/attendance")
def mark_attendance(lesson_id: int, body: AttendanceIn, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    if body.student_id not in group_member_ids(db, lesson.group_id):
        raise HTTPException(400, "Student is not in this group")
    row = db.scalar(select(Attendance).where(Attendance.lesson_id == lesson.id, Attendance.student_id == body.student_id))
    if row:
        row.status = body.status
    else:
        db.add(Attendance(lesson_id=lesson.id, student_id=body.student_id, status=body.status))
    db.commit()
    return {"ok": True}


@router.post("/lessons/{lesson_id}/attendance/all-present")
def all_present(lesson_id: int, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    existing = {a.student_id: a for a in db.scalars(select(Attendance).where(Attendance.lesson_id == lesson.id))}
    for sid in group_member_ids(db, lesson.group_id):
        if sid in existing:
            existing[sid].status = "P"
        else:
            db.add(Attendance(lesson_id=lesson.id, student_id=sid, status="P"))
    db.commit()
    return {"ok": True}


class TipIn(BaseModel):
    grades: list[int]
    text: str = Field(min_length=3, max_length=1000)
    send_at: datetime | None = None


@router.post("/tips")
def send_tip(body: TipIn, user: TeacherUser, db: DB) -> dict:
    taught = set()
    for g in teacher_svc.teacher_groups(db, user.id):
        taught |= {10, 11, 12} if int(g.grade_key) >= 10 else {int(g.grade_key)}
    grades = sorted(set(body.grades) & taught)
    if not grades:
        raise HTTPException(400, "Pick at least one of your grades")
    now = clock.now()
    send_at = body.send_at if body.send_at and body.send_at > now else now
    track = db.scalar(select(Track.id).where(Track.teacher_id == user.id))
    db.add(Tip(teacher_id=user.id, track_id=track, text=body.text.strip(), grades=grades, created_at=now,
               send_at=send_at))
    n = db.scalar(select(func.count()).select_from(StudentProfile).where(StudentProfile.grade.in_(grades))) or 0
    label = ", ".join(f"Grade {g}" for g in grades) if max(grades) < 10 else "Grades 10–12"
    events.log(db, user.id, f"{'Sent' if send_at == now else 'Scheduled'} a tip to {label}")
    db.commit()
    return {"count": n, "scheduled": send_at != now,
            "when": f"{clock.DAYS_SHORT[send_at.weekday()]} {clock.fmt_time(send_at.hour * 60 + send_at.minute)}"}


class AnswerIn(BaseModel):
    answer: str = Field(min_length=1, max_length=2000)


@router.post("/questions/{question_id}/answer")
def answer_question(question_id: int, body: AnswerIn, user: TeacherUser, db: DB) -> dict:
    q = db.get(Question, question_id)
    if not q or q.teacher_id != user.id:
        raise HTTPException(404, "Question not found")
    q.answer, q.answered_at = body.answer.strip(), clock.now()
    events.notify(db, q.student_id, f"{user.formal_short} answered your question: “{q.answer[:80]}”",
                  link="/student")
    db.commit()
    return {"ok": True}


class TextIn(BaseModel):
    text: str = Field(min_length=1, max_length=4000)


@router.post("/students/{student_id}/notes")
def add_note(student_id: int, body: TextIn, user: TeacherUser, db: DB) -> dict:
    own_student(db, user, student_id)
    db.add(Note(student_id=student_id, author_id=user.id, text=body.text.strip(), created_at=clock.now()))
    db.commit()
    return {"ok": True}


@router.post("/students/{student_id}/parent-message")
def parent_message(student_id: int, body: TextIn, user: TeacherUser, db: DB) -> dict:
    p = own_student(db, user, student_id)
    events.message_parent(db, student_id, user, body.text.strip())
    st = load_stats(db, [student_id]).get(student_id)
    name = db.get(User, student_id).name
    verb = "Flagged" if st and st.flags() else "Messaged"
    events.log(db, user.id, f"{verb} {name} and messaged their parent" if verb == "Flagged"
               else f"Messaged {name}’s parent")
    db.commit()
    return {"ok": True, "parent": p.parent_name}


# ------------------------------------------------------------------ planner
def lesson_time(db, planner: Planner, lesson: Lesson) -> LessonTime | None:
    return planner.lesson_time(db.get(Group, lesson.group_id), lesson.week)


def material_row(db, m: Material, total: int) -> dict:
    if m.is_slides:
        line = f"{m.name} · visible to students"
    elif m.kind == "VID":
        line = f"{m.name} · {len(m.viewed_by or [])} of {total} watched"
    else:
        line = f"{m.name} · {m.teacher_note}" if m.teacher_note else f"{m.name} · {len(m.viewed_by or [])} of {total} opened"
    return {"id": m.id, "kind": m.kind, "name": m.name, "line": line, "meta": m.meta, "is_slides": m.is_slides,
            "has_file": bool(m.file_path)}


@router.get("/planner")
def planner_view(user: TeacherUser, db: DB, group: int | None = None, week: int | None = None) -> dict:
    planner = Planner(db)
    groups = teacher_svc.teacher_groups(db, user.id)
    if not groups:
        raise HTTPException(404, "You have no groups yet")
    nxt = lesson_svc.next_lesson(db, planner, [g.id for g in groups])
    if group:
        g = own_group(db, user, group)
    else:
        g = db.get(Group, nxt[0].group_id) if nxt else groups[0]
    weeks = lesson_svc.module_weeks(db, g)
    now = clock.now()
    # The group's next lesson within the module
    next_week = None
    for w in weeks:
        t = planner.lesson_time(g, w)
        if t and t.ends_at > now:
            next_week = w
            break
    sel_week = week if week in weeks else (next_week or weeks[-1])
    strip = []
    for i, w in enumerate(weeks):
        lesson = lesson_svc.get_or_create(db, g, w)
        t = planner.lesson_time(g, w)
        state, label = lesson_svc.status_label(t, w == next_week, bool(lesson.published_at))
        strip.append({"week": w, "n": i + 1, "t": lesson.title, "state": state, "label": label, "lesson_id": lesson.id})
    lesson = lesson_svc.get_or_create(db, g, sel_week)
    t = planner.lesson_time(g, sel_week)
    lv, track = level_of(db, g)
    level_skills = list(db.scalars(select(Skill).where(Skill.level_id == g.level_id).order_by(Skill.position))) if lv else []
    skill_rows = []
    for sid in lesson.skill_ids or []:
        s = db.get(Skill, sid)
        if s:
            slv = db.get(Level, s.level_id)
            strack = db.get(Track, slv.track_id)
            skill_rows.append({"id": s.id, "name": f"{s.name} · {strack.short} L{slv.number}"})
    members = group_member_ids(db, g.id)
    hw = lesson_svc.homework_of(db, lesson)
    homework = None
    if hw:
        submitted = sub_counts(db, hw.id)
        opened = hw.published_at and hw.published_at.date() == clock.today()
        homework = {
            "id": hw.id, "title": hw.title, "description": hw.description, "due_at": hw.due_at.isoformat(timespec="minutes"),
            "due_label": f"Due {clock.fmt_day(hw.due_at.date())}", "points": hw.points, "rubric": hw.rubric,
            "response_type": hw.response_type, "published": bool(hw.published_at),
            "status_line": f"{submitted} of {len(members)} submitted" + (" · opens today" if opened else
                                                                         (" · draft" if not hw.published_at else "")),
        }
    lib = []
    prev = None
    for item in db.scalars(select(LibraryItem).order_by(LibraryItem.id)):
        if item.kind == "previous":
            if item.level_id and item.level_id == g.level_id and item.shared_by_id == user.id:
                prev = {"id": item.id, "title": item.title, "meta": item.meta, "blocks": item.blocks}
            continue
        if item.track_id and track and item.track_id != track.id and user.id not in (item.used_by or []):
            continue
        sharer = db.get(User, item.shared_by_id)
        used = len(item.used_by or [])
        target = min(sel_week + 1, weeks[-1]) if sel_week < weeks[-1] else sel_week
        lib.append({"id": item.id, "title": item.title, "target_week": target,
                    "meta": f"Shared by {sharer.formal_short} · used by {used} teacher{'s' if used != 1 else ''}",
                    "used": user.id in (item.used_by or [])})
    on = autos.settings(db, user.id)
    db.commit()
    return {
        "groups": [{"id": x.id, "name": x.name} for x in groups], "group_id": g.id,
        "module_label": f"Module · {short(g)}",
        "weeks": strip,
        "lesson": {
            "id": lesson.id, "week": sel_week, "n": weeks.index(sel_week) + 1, "of": len(weeks), "title": lesson.title,
            "time_label": f"{clock.fmt_day(t.day)} · {t.label} · {g.room}" if t else "Not on the schedule this week",
            "minutes": t.minutes if t else 0, "published": bool(lesson.published_at),
            "updated_after_publish": bool(lesson.published_at and lesson.updated_at and lesson.updated_at > lesson.published_at),
            "blocks": lesson.blocks or [], "skills": skill_rows, "extra_skills": lesson.extra_skills or [],
            "available_skills": [{"id": s.id, "name": s.name} for s in level_skills],
            "student_prep": lesson.student_prep or [], "slides": len(lesson.slides or []),
            "teacher": user.formal, "group_name": g.name, "room": g.room,
        },
        "homework": homework,
        "materials": [material_row(db, m, len(members)) for m in lesson.materials],
        "automations": [{"key": key, "t": tt, "d": d, "on": on[key]} for key, tt, d in autos.AUTOMATIONS],
        "library": lib, "previous": prev, "size": len(members),
    }


class LessonIn(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=160)
    blocks: list[dict] | None = None
    skill_ids: list[int] | None = None
    extra_skills: list[str] | None = None
    student_prep: list[dict] | None = None


def clean_blocks(blocks: list[dict]) -> list[dict]:
    out = []
    for b in blocks:
        kind = b.get("kind") if b.get("kind") in BLOCK_KINDS else "teach"
        try:
            minutes = max(1, min(180, int(b.get("min", 10))))
        except (TypeError, ValueError):
            minutes = 10
        row = {"kind": kind, "min": minutes, "t": str(b.get("t", "")).strip()[:120] or "Untitled block",
               "d": str(b.get("d", "")).strip()[:200]}
        for extra in ("desc", "badge"):
            if b.get(extra):
                row[extra] = str(b[extra])[:400]
        out.append(row)
    return out


@router.put("/lessons/{lesson_id}")
def update_lesson(lesson_id: int, body: LessonIn, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    if body.title is not None:
        lesson.title = body.title.strip()
    if body.blocks is not None:
        lesson.blocks = clean_blocks(body.blocks)
    if body.skill_ids is not None:
        lesson.skill_ids = [s for s in body.skill_ids if db.get(Skill, s)]
    if body.extra_skills is not None:
        lesson.extra_skills = [s.strip()[:40] for s in body.extra_skills if s.strip()]
    if body.student_prep is not None:
        lesson.student_prep = body.student_prep
    lesson.updated_at = clock.now()
    db.commit()
    return {"ok": True}


@router.post("/lessons/{lesson_id}/publish")
def publish_lesson(lesson_id: int, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    g = db.get(Group, lesson.group_id)
    first = lesson.published_at is None
    lesson.published_at = lesson.updated_at = clock.now()
    week_n = lesson_svc.module_week_number(db, g, lesson.week)
    for sid in group_member_ids(db, g.id):
        events.notify(db, sid, f"{user.formal_short} {'published' if first else 'updated'} week {week_n}: {lesson.title}",
                      link=f"/student/session/{lesson.id}")
    events.log(db, user.id, f"Published the week {week_n} plan for {short(g)}")
    db.commit()
    return {"ok": True}


class HomeworkIn(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    description: str = ""
    due_at: datetime
    points: int = Field(default=20, ge=1, le=100)
    rubric: list[dict] = []
    response_type: str = "File or photo"
    publish: bool = True


@router.put("/lessons/{lesson_id}/homework")
def save_homework(lesson_id: int, body: HomeworkIn, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    g = db.get(Group, lesson.group_id)
    rubric = [{"c": str(r.get("c", "")).strip()[:120], "p": int(r.get("p", 0) or 0)}
              for r in body.rubric if str(r.get("c", "")).strip()]
    hw = lesson_svc.homework_of(db, lesson)
    now = clock.now()
    kind = "journal" if "journal" in body.title.lower() else "homework"
    if not hw:
        hw = Assignment(group_id=g.id, lesson_id=lesson.id, kind=kind, created_at=now, title=body.title, due_at=body.due_at)
        db.add(hw)
    newly = body.publish and not hw.published_at
    hw.title, hw.description, hw.due_at, hw.points = body.title.strip(), body.description.strip(), body.due_at, body.points
    hw.rubric, hw.response_type, hw.kind = rubric, body.response_type, kind
    if body.publish and not hw.published_at:
        hw.published_at = now
    db.flush()
    if newly:
        for sid in group_member_ids(db, g.id):
            events.notify(db, sid, f"{user.formal_short} published new homework: {hw.title} · due "
                                   f"{clock.fmt_day(hw.due_at.date())}", link=f"/student/session/{lesson.id}")
        events.log(db, user.id, f"Published week {lesson_svc.module_week_number(db, g, lesson.week)} homework for {short(g)}")
    db.commit()
    return {"id": hw.id}


@router.post("/lessons/{lesson_id}/materials")
async def upload_material(lesson_id: int, user: TeacherUser, db: DB, file: UploadFile = File(...)) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    rel, size = await save_upload(file, f"materials/{lesson.id}")
    kind = kind_for(file.filename or "")
    is_slides = kind == "PDF" and not lesson_svc.slides_material(lesson)
    name = f"Week {lesson_svc.module_week_number(db, db.get(Group, lesson.group_id), lesson.week)} slides" \
        if is_slides else (file.filename or "File").rsplit(".", 1)[0]
    m = Material(lesson_id=lesson.id, kind=kind if kind != "IMG" else "DOC", name=name,
                 meta=f"{size_label(size)}", is_slides=is_slides, file_path=rel, file_size=size,
                 uploaded_at=clock.now(), viewed_by=[])
    db.add(m)
    g = db.get(Group, lesson.group_id)
    events.log(db, user.id, f"Uploaded {'week ' + str(lesson_svc.module_week_number(db, g, lesson.week)) + ' slides' if is_slides else m.name} for {short(g)}")
    db.commit()
    return {"id": m.id, "slides": is_slides}


@router.delete("/materials/{material_id}")
def delete_material(material_id: int, user: TeacherUser, db: DB) -> dict:
    m = db.get(Material, material_id)
    if not m:
        raise HTTPException(404, "Material not found")
    own_lesson(db, user, m.lesson_id)
    db.delete(m)
    db.commit()
    return {"ok": True}


class ToggleIn(BaseModel):
    enabled: bool


@router.put("/automations/{key}")
def set_automation(key: str, body: ToggleIn, user: TeacherUser, db: DB) -> dict:
    if key not in {k for k, _, _ in autos.AUTOMATIONS}:
        raise HTTPException(404, "Unknown automation")
    autos.set_enabled(db, user.id, key, body.enabled)
    db.commit()
    return {"ok": True}


class UseIn(BaseModel):
    lesson_id: int


@router.post("/library/{item_id}/use")
def use_library(item_id: int, body: UseIn, user: TeacherUser, db: DB) -> dict:
    item = db.get(LibraryItem, item_id)
    if not item or item.kind != "kit":
        raise HTTPException(404, "Library item not found")
    lesson = own_lesson(db, user, body.lesson_id)
    lesson.blocks = [*(lesson.blocks or []), *item.blocks]
    lesson.updated_at = clock.now()
    if user.id not in (item.used_by or []):
        item.used_by = [*(item.used_by or []), user.id]
    db.commit()
    g = db.get(Group, lesson.group_id)
    return {"week": lesson_svc.module_week_number(db, g, lesson.week)}


@router.post("/lessons/{lesson_id}/copy-previous")
def copy_previous(lesson_id: int, user: TeacherUser, db: DB) -> dict:
    lesson = own_lesson(db, user, lesson_id)
    g = db.get(Group, lesson.group_id)
    prev = db.scalar(select(LibraryItem).where(LibraryItem.kind == "previous", LibraryItem.level_id == g.level_id,
                                               LibraryItem.shared_by_id == user.id))
    if not prev:
        raise HTTPException(404, "No plan from last year for this level")
    lesson.blocks = list(prev.blocks)
    lesson.updated_at = clock.now()
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- gradebook
def cell_state(a: Assignment, s: Submission | None, now: datetime) -> dict:
    if s and s.status == "graded":
        return {"state": "score", "score": round(s.score / a.points * 20) if a.points != 20 else s.score, "sub_id": s.id}
    if s and s.status == "submitted":
        return {"state": "todo", "sub_id": s.id}
    if a.due_at > now:
        return {"state": "due", "due": a.due_at.day, "sub_id": s.id if s else None}
    return {"state": "missing"}


def gradebook_group(db, user: TeacherUser, group: int | None, planner: Planner) -> Group:
    if group:
        return own_group(db, user, group)
    groups = teacher_svc.teacher_groups(db, user.id)
    if not groups:
        raise HTTPException(404, "You have no groups yet")
    nxt = lesson_svc.next_lesson(db, planner, [g.id for g in groups])
    return db.get(Group, nxt[0].group_id) if nxt else groups[0]


def group_assignments(db, g: Group) -> list[Assignment]:
    return list(db.scalars(select(Assignment).where(Assignment.group_id == g.id, Assignment.published_at.is_not(None))
                           .order_by(Assignment.due_at)))


@router.get("/gradebook")
def gradebook(user: TeacherUser, db: DB, group: int | None = None) -> dict:
    planner = Planner(db)
    g = gradebook_group(db, user, group, planner)
    now = clock.now()
    ids = group_member_ids(db, g.id)
    local = load_stats(db, ids, planner, group_ids=[g.id])
    overall = load_stats(db, ids, planner)
    assignments = group_assignments(db, g)
    subs = {(s.assignment_id, s.student_id): s for s in db.scalars(
        select(Submission).where(Submission.assignment_id.in_([a.id for a in assignments] or [0])))}
    lessons = {lesson.id: lesson for lesson in db.scalars(select(Lesson).where(Lesson.group_id == g.id))}
    att = {(a.lesson_id, a.student_id): a.status for a in db.scalars(
        select(Attendance).where(Attendance.lesson_id.in_(list(lessons) or [0])))}
    columns = []
    for a in assignments:
        wk = lessons[a.lesson_id].week if a.lesson_id in lessons else clock.week_of(a.due_at.date())
        columns.append({"assignment_id": a.id, "label": f"W{lesson_svc.module_week_number(db, g, wk)}",
                        "title": a.title, "kind": a.kind, "due": f"Due {clock.fmt_day(a.due_at.date())}",
                        "lesson_id": a.lesson_id})
    lv, track = level_of(db, g)
    first_skill = db.scalar(select(Skill).where(Skill.level_id == g.level_id, Skill.position == 0)) if lv else None
    mastered = set(db.scalars(select(StudentSkill.student_id).where(
        StudentSkill.skill_id == first_skill.id, StudentSkill.status == "done"))) if first_skill else set()

    rows = []
    for sid in ids:
        u, p = db.get(User, sid), db.get(StudentProfile, sid)
        cells, total, counted = [], 0.0, 0
        for a in assignments:
            s = subs.get((a.id, sid))
            c = cell_state(a, s, now)
            st = att.get((a.lesson_id, sid))
            c["att"] = {"L": "late", "A": "absent"}.get(st or "", "")
            cells.append(c)
            if c["state"] == "score":
                total += c["score"]
                counted += 1
            elif c["state"] == "missing":
                counted += 1
        rows.append({"student_id": sid, "name": u.name, "cls": p.class_name,
                     "flag": bool(overall[sid].flags()), "quiz": local[sid].quiz_avg,
                     "avg": round(total / counted, 1) if counted else None, "cells": cells})
    queue = []
    for a in assignments:
        for sid in ids:
            s = subs.get((a.id, sid))
            if s and s.status == "submitted":
                queue.append({"submission_id": s.id, "assignment_id": a.id, "student": db.get(User, sid).name,
                              "title": a.title, "submitted": clock.fmt_day(s.submitted_at.date()) if s.submitted_at else "",
                              "late": bool(s.submitted_at and s.submitted_at > a.due_at)})
    cells_all = [c for st in local.values() for c in st.cells]
    att_pct = round(sum(1 for c in cells_all if c.status in ("P", "L")) / len(cells_all) * 100) if cells_all else None
    hw_avgs = [st.homework_avg for st in local.values() if st.homework_avg is not None]
    quiz_avgs = [st.quiz_avg for st in local.values() if st.quiz_avg is not None]
    delivered = sorted({lessons[lid].week for (lid, _) in att})
    span = f"weeks {lesson_svc.module_week_number(db, g, delivered[0])}–" \
           f"{lesson_svc.module_week_number(db, g, max([lessons[a.lesson_id].week for a in assignments if a.lesson_id in lessons] or delivered))}" \
        if delivered else "no sessions yet"
    level_label = f"{track.name} Level {lv.number}" if lv else (g.spec_track.name if g.spec_track_id else g.name)
    return {
        "groups": [{"id": x.id, "name": x.name} for x in teacher_svc.teacher_groups(db, user.id)],
        "group_id": g.id,
        "title": f"{short(g)} · {level_label} · {len(ids)} students · {span}",
        "kpis": {
            "attendance": att_pct,
            "hw_avg": round(sum(hw_avgs) / len(hw_avgs), 1) if hw_avgs else None,
            "quiz_avg": round(sum(quiz_avgs) / len(quiz_avgs)) if quiz_avgs else None,
            "skill": first_skill.name if first_skill else None,
            "skill_done": len(mastered & set(ids)), "n": len(ids),
        },
        "columns": columns, "rows": rows, "queue": queue,
    }


@router.get("/gradebook/students/{student_id}")
def gradebook_student(student_id: int, user: TeacherUser, db: DB, group: int) -> dict:
    planner = Planner(db)
    g = own_group(db, user, group)
    if student_id not in group_member_ids(db, g.id):
        raise HTTPException(404, "Student is not in this group")
    u, p = db.get(User, student_id), db.get(StudentProfile, student_id)
    local = load_stats(db, [student_id], planner, group_ids=[g.id])[student_id]
    statuses = list(db.scalars(select(Attendance.status).join(Lesson, Lesson.id == Attendance.lesson_id)
                               .where(Lesson.group_id == g.id, Attendance.student_id == student_id)))
    lv, track = level_of(db, g)
    skills = []
    if lv:
        ss = dict(db.execute(select(StudentSkill.skill_id, StudentSkill.status).where(StudentSkill.student_id == student_id)).all())
        for s in db.scalars(select(Skill).where(Skill.level_id == lv.id).order_by(Skill.position)):
            st = ss.get(s.id, "next")
            if st == "done":
                skills.append({"n": s.name, "st": "Mastered", "tone": "ok"})
            elif st == "now":
                skills.append({"n": s.name, "st": "Needs review" if s.position == 0 else "In progress",
                               "tone": "warn" if s.position == 0 else "info"})
            else:
                skills.append({"n": s.name, "st": "Coming up", "tone": "muted"})
    rec = None
    spec_id = REC_TRACK.get(track.id) if track else None
    if spec_id:
        spec = db.get(SpecTrack, spec_id)
        on = bool(db.scalar(select(Recommendation.id).where(
            Recommendation.teacher_id == user.id, Recommendation.student_id == student_id,
            Recommendation.spec_track_id == spec_id)))
        rec = {"track_id": spec_id, "name": spec.name, "on": on, "grade": p.grade}
    notes = [{"who": db.get(User, n.author_id).name, "text": n.text, "when": clock.fmt_day(n.created_at.date())}
             for n in db.scalars(select(Note).where(Note.student_id == student_id).order_by(Note.created_at.desc()).limit(5))]
    since = max(7, min(p.grade, p.grade - (clock.today().year - p.joined_on.year) + (0 if clock.today().month >= 9 else 1)))
    return {
        "id": student_id, "name": u.name, "first": u.first_name, "ini": u.initials, "cls": p.class_name,
        "since": f"in Ithraa since Grade {since}", "parent": p.parent_name,
        "attend": f"{sum(1 for s in statuses if s in ('P', 'L'))}/{len(statuses)}",
        "avg": round(local.homework_avg, 1) if local.homework_avg is not None else None,
        "quiz": local.quiz_avg, "skills": skills, "rec": rec, "notes": notes,
        "flags": load_stats(db, [student_id], planner)[student_id].flags(),
    }


class GradeIn(BaseModel):
    score: int = Field(ge=0)
    feedback: str = ""
    portfolio: bool = False


@router.get("/submissions/{submission_id}")
def submission_detail(submission_id: int, user: TeacherUser, db: DB) -> dict:
    s = db.get(Submission, submission_id)
    if not s:
        raise HTTPException(404, "Submission not found")
    a = db.get(Assignment, s.assignment_id)
    own_group(db, user, a.group_id)
    return {"id": s.id, "student": db.get(User, s.student_id).name, "title": a.title, "points": a.points,
            "rubric": a.rubric, "file_name": s.file_name, "has_file": bool(s.file_path),
            "submitted": clock.fmt_day(s.submitted_at.date()) if s.submitted_at else None,
            "late": bool(s.submitted_at and s.submitted_at > a.due_at), "score": s.score, "feedback": s.feedback or "",
            "status": s.status}


@router.post("/submissions/{submission_id}/grade")
def grade(submission_id: int, body: GradeIn, user: TeacherUser, db: DB) -> dict:
    s = db.get(Submission, submission_id)
    if not s:
        raise HTTPException(404, "Submission not found")
    a = db.get(Assignment, s.assignment_id)
    g = own_group(db, user, a.group_id)
    if s.status not in ("submitted", "graded"):
        raise HTTPException(400, "Nothing has been handed in yet")
    if body.score > a.points:
        raise HTTPException(400, f"Score can be at most {a.points}")
    first = s.status != "graded"
    s.status, s.score, s.feedback = "graded", body.score, body.feedback.strip() or None
    s.graded_at, s.graded_by = clock.now(), user.id
    student = db.get(User, s.student_id)
    events.notify(db, s.student_id, f"Your {a.title} was graded · {body.score}/{a.points}",
                  link=f"/student/session/{a.lesson_id}" if a.lesson_id else "/student")
    if first:
        events.log(db, s.student_id, f"{a.title} graded · {body.score}/{a.points}")
        events.log(db, user.id, f"Graded {student.name}’s {a.kind} for {short(g)}")
    if a.lesson_id and body.score / a.points >= 0.7:
        for skill_id in db.get(Lesson, a.lesson_id).skill_ids or []:
            set_skill(db, s.student_id, skill_id, "done")
    added = False
    if body.portfolio and not db.scalar(select(PortfolioItem.id).where(PortfolioItem.submission_id == s.id)):
        lv, track = level_of(db, g)
        track_id = track.id if track else SPEC_TO_TRACK.get(g.spec_track_id or "", "coding")
        db.add(PortfolioItem(student_id=s.student_id, track_id=track_id, level_id=lv.id if lv else None,
                             submission_id=s.id, title=a.title, kind="Journal" if a.kind == "journal" else "Homework",
                             made_on=(s.submitted_at or clock.now()).date(), score=f"{body.score}/{a.points}",
                             feedback=body.feedback.strip(), teacher_id=user.id, featured=False,
                             skills=[db.get(Skill, x).name for x in (db.get(Lesson, a.lesson_id).skill_ids if a.lesson_id else []) if db.get(Skill, x)],
                             art=[random.randint(48, 64), random.randint(22, 40)]))
        events.notify(db, s.student_id, f"{user.formal_short} added “{a.title}” to your portfolio", link="/student/portfolio")
        added = True
    db.commit()
    return {"ok": True, "portfolio": added}


class RecIn(BaseModel):
    track_id: str


@router.post("/students/{student_id}/recommend")
def recommend(student_id: int, body: RecIn, user: TeacherUser, db: DB) -> dict:
    own_student(db, user, student_id)
    if not db.get(SpecTrack, body.track_id):
        raise HTTPException(404, "Unknown track")
    row = db.scalar(select(Recommendation).where(Recommendation.teacher_id == user.id,
                                                 Recommendation.student_id == student_id,
                                                 Recommendation.spec_track_id == body.track_id))
    if row:
        db.delete(row)
        on = False
    else:
        db.add(Recommendation(teacher_id=user.id, student_id=student_id, spec_track_id=body.track_id,
                              created_at=clock.now()))
        on = True
        track = db.get(SpecTrack, body.track_id)
        events.notify(db, student_id, f"{user.formal_short} recommended you for {track.name}",
                      link="/student/specialization", ref=f"rec:{user.id}:{body.track_id}:{student_id}")
    db.commit()
    return {"on": on}


@router.get("/gradebook/export.xlsx")
def export_xlsx(user: TeacherUser, db: DB, group: int):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = gradebook(user, db, group)
    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    ws.append([data["title"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    head = ["Student", "Class", *[f"{c['label']} · {c['title']}" for c in data["columns"]], "Quizzes %", "Average /20", "Check in"]
    ws.append(head)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    fills = {"score_hi": "DCF4EA", "score_mid": "F1F2F8", "score_lo": "FFF1DE", "missing": "FFE6E8"}
    for r in data["rows"]:
        vals = []
        for c in r["cells"]:
            vals.append(c["score"] if c["state"] == "score" else {"missing": "Missing", "todo": "To grade",
                                                                    "due": "Not due"}.get(c["state"], ""))
        ws.append([r["name"], r["cls"], *vals, r["quiz"], r["avg"], "Yes" if r["flag"] else ""])
        row = ws.max_row
        for i, c in enumerate(r["cells"]):
            key = "missing" if c["state"] == "missing" else (
                ("score_hi" if c["score"] >= 17 else "score_mid" if c["score"] >= 14 else "score_lo")
                if c["state"] == "score" else None)
            if key:
                ws.cell(row=row, column=3 + i).fill = PatternFill("solid", fgColor=fills[key])
    ws.column_dimensions["A"].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    g = db.get(Group, group)
    name = f"gradebook-{g.name.replace(' · ', '-').replace(' ', '_')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.post("/gradebook/progress")
def send_progress(user: TeacherUser, db: DB, group: int) -> dict:
    planner = Planner(db)
    g = own_group(db, user, group)
    ids = group_member_ids(db, g.id)
    local = load_stats(db, ids, planner, group_ids=[g.id])
    for sid in ids:
        st = local[sid]
        u, p = db.get(User, sid), db.get(StudentProfile, sid)
        parts = []
        if st.attendance_pct is not None:
            parts.append(f"attendance {st.attendance_pct}%")
        if st.homework_avg is not None:
            parts.append(f"homework average {st.homework_avg:.1f}/20")
        if st.quiz_avg is not None:
            parts.append(f"quiz average {st.quiz_avg}%")
        missing = sum(1 for h in st.homework if h.missing)
        if missing:
            parts.append(f"{missing} missing homework{'s' if missing > 1 else ''}")
        text = (f"Dear {p.parent_name}, here is {u.first_name}’s progress in {g.name} so far: "
                f"{', '.join(parts) or 'no graded work yet'}. — {user.formal}")
        events.message_parent(db, sid, user, text, kind="progress")
    events.log(db, user.id, f"Sent progress to {len(ids)} {short(g)} parents")
    db.commit()
    return {"count": len(ids)}
